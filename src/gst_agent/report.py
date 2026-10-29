"""Turn the reconciliation result into a colour-coded Excel report.

Sheets:
  Summary          counts per status + tax at risk on invoices missing from 2B
  All invoices     every invoice in your books, colour-coded by status
  Needs follow-up  only the problem invoices (what to chase suppliers about)
  Extra in 2B      invoices on the portal that are not in your books
  Needs review     invoices the AI could not read reliably (set aside, not matched)

Run from the repo root (after the matcher):
    python -m src.gst_agent.report
"""
import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

STATUS_LABEL = {
    "matched": "Matched",
    "amount_mismatch": "Amount mismatch",
    "invoice_no_format_diff": "Invoice no. written differently",
    "missing_in_2b": "Missing in 2B",
}
STATUS_FILL = {
    "matched": "C6EFCE",              # green
    "amount_mismatch": "FFC7CE",      # red
    "invoice_no_format_diff": "FFEB9C",  # yellow
    "missing_in_2b": "F8CBAD",        # orange
}
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def explain(row: dict) -> str:
    """One plain-English line saying what is wrong with this invoice."""
    s, d = row["status"], row["details"]
    if s == "matched":
        return ""
    if s == "missing_in_2b":
        return "Not found on the portal. Ask the supplier to file it."
    if s == "invoice_no_format_diff":
        return f"Books: {d['books']} / Portal: {d['portal']}. Probably the same invoice."
    parts = [f"{f}: books {v['books']:,.2f} vs portal {v['portal']:,.2f}" for f, v in d.items()]
    return "; ".join(parts)


def write_table(ws, headers: list[str], rows: list[list], fills: list[str] | None = None) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for i, r in enumerate(rows):
        ws.append(r)
        if fills:
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor=fills[i])
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for col in range(1, len(headers) + 1):
        longest = max(len(str(c.value or "")) for c in ws[get_column_letter(col)])
        ws.column_dimensions[get_column_letter(col)].width = min(max(longest + 2, 12), 70)


def tax_of(book: dict) -> float:
    return float(book["cgst"]) + float(book["sgst"]) + float(book["igst"])


def build_report(recon: dict, books_by_no: dict, out_path: Path) -> None:
    results, extras = recon["results"], recon["extra_in_2b"]
    wb = Workbook()

    # --- Summary ---
    ws = wb.active
    ws.title = "Summary"
    counts = {k: 0 for k in STATUS_LABEL}
    for r in results:
        counts[r["status"]] += 1
    at_risk = sum(tax_of(books_by_no[r["invoice_no"]]) for r in results
                  if r["status"] == "missing_in_2b" and r["invoice_no"] in books_by_no)
    ws.append(["GST reconciliation summary"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Invoices in books", len(results)])
    for key, label in STATUS_LABEL.items():
        ws.append([label, counts[key]])
        ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=STATUS_FILL[key])
    review = recon.get("needs_review", [])
    ws.append(["Extra in 2B (not in books)" + (" - may include set-aside invoices" if review else ""), len(extras)])
    ws.append(["Set aside for review (not matched)", len(review)])
    ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor="BDD7EE")
    ws.append([])
    ws.append(["Tax on invoices missing in 2B (input credit possibly at risk)", round(at_risk, 2)])
    ws.cell(ws.max_row, 2).number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 62
    ws.column_dimensions["B"].width = 16

    # --- All invoices ---
    headers = ["Invoice no.", "Supplier", "GSTIN", "Status", "What is wrong"]
    ws_all = wb.create_sheet("All invoices")
    write_table(ws_all, headers,
                [[r["invoice_no"], r["supplier_name"], r["supplier_gstin"],
                  STATUS_LABEL[r["status"]], explain(r)] for r in results],
                [STATUS_FILL[r["status"]] for r in results])

    # --- Needs follow-up ---
    problems = [r for r in results if r["status"] != "matched"]
    ws_fu = wb.create_sheet("Needs follow-up")
    write_table(ws_fu, headers,
                [[r["invoice_no"], r["supplier_name"], r["supplier_gstin"],
                  STATUS_LABEL[r["status"]], explain(r)] for r in problems],
                [STATUS_FILL[r["status"]] for r in problems])

    # --- Extra in 2B ---
    ws_x = wb.create_sheet("Extra in 2B")
    write_table(ws_x, ["Invoice no.", "GSTIN", "Date", "Taxable value", "CGST", "SGST", "IGST"],
                [[p["invoice_no"], p["supplier_gstin"], p["invoice_date"], p["taxable_value"],
                  p["cgst"], p["sgst"], p["igst"]] for p in extras])

    # --- Needs review ---
    if review:
        ws_r = wb.create_sheet("Needs review")
        write_table(ws_r, ["File", "Supplier (as read)", "Invoice no. (as read)", "Why it was set aside"],
                    [[r["source"], r["supplier_name"], r["invoice_no"], "; ".join(r["reasons"])] for r in review],
                    ["BDD7EE"] * len(review))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--recon", default="data/recon_result.json")
    ap.add_argument("--extracted", default="data/extracted")
    ap.add_argument("--out", default="output/gst_reconciliation_report.xlsx")
    args = ap.parse_args()

    recon = json.loads(Path(args.recon).read_text())
    books = [json.loads(p.read_text()) for p in sorted(Path(args.extracted).glob("*.json"))]
    build_report(recon, {b["invoice_no"]: b for b in books}, Path(args.out))
    print(f"Report saved to {args.out}")


if __name__ == "__main__":
    main()
